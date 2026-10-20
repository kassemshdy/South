"""The administrator's Excel export (``app.services.data_export``).

It exists so the people who applied, and what they listed, survive losing the
platform. These pin the three things that make it fit for that and safe to
hand out: it is complete and linked by id, it never carries a credential, and
nothing a member of the public typed can run as a formula when it is opened.
"""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.models.taxonomy import Category, Location
from app.models.user import User
from tests.conftest import admin_headers, make_owner, sign_in
from tests.samples import ar

FORMULA = '=HYPERLINK("https://example.com","x")'


def _export(client: TestClient):
    response = client.get("/api/admin/export", headers=admin_headers(client))
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["cache-control"] == "no-store"
    assert "attachment" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content))


def _rows(sheet) -> list[dict[str, object]]:
    header, *rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(header, row, strict=True)) for row in rows]


def test_the_export_holds_people_and_what_they_listed_linked_by_id(
    client: TestClient, admin: User, category: Category, location: Location
) -> None:
    owner = make_owner(
        "03981001",
        full_name=ar("business.applicant"),
        birth_year=1990,
        personal_phone_number="+96170000001",
    )
    business = client.post(
        "/api/businesses",
        headers=sign_in(client, "03981001"),
        json={
            "name": ar("business.first"),
            "category_id": str(category.id),
            "location_id": str(location.id),
            "latitude": 33.27,
            "longitude": 35.2,
        },
    )
    assert business.status_code in (200, 201), business.text

    workbook = _export(client)
    assert workbook.sheetnames[:4] == ["people", "talents", "businesses", "items"]
    assert {"categories", "locations", "talent_skills"} <= set(workbook.sheetnames)

    people = {row["id"]: row for row in _rows(workbook["people"])}
    person = people[str(owner.id)]
    assert person["phone_number"] == owner.phone_number
    assert person["full_name"] == ar("business.applicant")
    assert person["birth_year"] == 1990
    assert person["personal_phone_number"] == "+96170000001"

    businesses = _rows(workbook["businesses"])
    assert [row["owner_id"] for row in businesses] == [str(owner.id)]
    assert businesses[0]["name"] == ar("business.first")
    assert businesses[0]["status"] == "DRAFT"
    category_ids = {row["id"] for row in _rows(workbook["categories"])}
    assert businesses[0]["category_id"] in category_ids


def test_the_export_never_carries_a_credential(client: TestClient, admin: User) -> None:
    make_owner("03981002")
    workbook = _export(client)

    header = next(workbook["people"].iter_rows(values_only=True))
    assert "password_hash" not in header
    assert "token_version" not in header
    # Nor anywhere else, under another name: a bcrypt hash is recognisable.
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            assert not any(isinstance(value, str) and value.startswith("$2") for value in row)


def test_text_that_looks_like_a_formula_stays_text(
    client: TestClient, admin: User
) -> None:
    make_owner("03981003", full_name=FORMULA, residence_place="tab\there\x0bvt")
    workbook = load_workbook(
        BytesIO(client.get("/api/admin/export", headers=admin_headers(client)).content)
    )

    sheet = workbook["people"]
    header = next(sheet.iter_rows(values_only=True))
    name_column = header.index("full_name") + 1
    cells = [row[name_column - 1] for row in sheet.iter_rows(min_row=2)]
    formula_cells = [cell for cell in cells if cell.value == FORMULA]
    assert formula_cells, "the formula-like name was not exported as its text"
    assert all(cell.data_type == "s" for cell in formula_cells)

    # A control character a pasted value can carry is dropped rather than
    # failing the whole export; tab survives.
    residence = [row["residence_place"] for row in _rows(sheet)]
    assert "tab\therevt" in residence


def test_only_an_administrator_can_export(client: TestClient, admin: User) -> None:
    assert client.get("/api/admin/export").status_code == 401
    owner = client.get("/api/admin/export", headers=sign_in(client, "03981004"))
    assert owner.status_code == 403
